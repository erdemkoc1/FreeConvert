import os
import tempfile
import unittest

from reportlab.pdfgen import canvas
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from converter.pdf_to_word import convert_pdf_to_word


class PdfToWordTableTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(
            prefix="dc_table_tests_")
        self.root = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def _convert(self, source, name):
        output = os.path.join(self.root, name + ".docx")
        ok, info = convert_pdf_to_word(source, output, mode="sabit")
        self.assertTrue(ok, info)
        return Document(output)

    def _ruled_pdf(self, name, data, spans=()):
        path = os.path.join(self.root, name + ".pdf")
        table = Table(data, colWidths=[60, 150, 55, 70])
        style = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ]
        style.extend(("SPAN", *span) for span in spans)
        table.setStyle(TableStyle(style))
        SimpleDocTemplate(path, pagesize=A4).build([table])
        return path

    def test_ruled_table_cells_and_dimensions(self):
        source = self._ruled_pdf("ruled", [
            ["Code", "Description", "Qty", "Price"],
            ["A", "Apple", "2", "10.00"],
            ["B", "Orange", "4", "12.50"],
        ])
        document = self._convert(source, "ruled")
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (3, 4))
        self.assertEqual(table.cell(1, 1).text, "Apple")
        self.assertEqual(table.cell(2, 3).text, "12.50")

    def test_unruled_table_is_detected(self):
        path = os.path.join(self.root, "unruled.pdf")
        pdf = canvas.Canvas(path, pagesize=A4)
        rows = [
            ("Name", "Age", "City"),
            ("Ali", "30", "Ankara"),
            ("Veli", "25", "Izmir"),
            ("Deli", "40", "Bursa"),
        ]
        y = 760
        for row in rows:
            for x, value in zip((80, 220, 330), row):
                pdf.setFont("Helvetica", 11)
                pdf.drawString(x, y, value)
            y -= 28
        pdf.save()
        converted = self._convert(path, "unruled")
        self.assertEqual(len(converted.tables), 1)
        self.assertEqual(
            (len(converted.tables[0].rows), len(converted.tables[0].columns)),
            (4, 3),
        )

    def test_colspan_and_rowspan_are_emitted(self):
        source = self._ruled_pdf(
            "spans",
            [
                ["Description", "", "", ""],
                ["Code", "Name", "Qty", "Price"],
                ["A", "Apple", "2", "10.00"],
                ["B", "Orange", "4", "12.50"],
            ],
            spans=[((0, 0), (3, 0))],
        )
        document = self._convert(source, "spans")
        table = document.tables[0]
        self.assertEqual(table.cell(0, 0).text, "Description")
        self.assertIn("<w:gridSpan", table._tbl.xml)

        source = self._ruled_pdf(
            "rowspan",
            [
                ["Group", "Value"],
                ["A", "one"],
                ["", "two"],
                ["", "three"],
            ],
            spans=[((0, 1), (0, 3))],
        )
        document = self._convert(source, "rowspan")
        table = document.tables[0]
        self.assertEqual(table.cell(1, 0).text, "A")
        self.assertIn("<w:vMerge", table._tbl.xml)

    def test_multiline_cell_and_prose_negative(self):
        path = os.path.join(self.root, "multiline.pdf")
        table = Table([
            ["Description", "Qty"],
            ["Line one\nLine two\nCity", "2"],
            ["Other", "3"],
        ], colWidths=[220, 60], rowHeights=[25, 70, 25])
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ]))
        SimpleDocTemplate(path, pagesize=A4).build([table])
        document = self._convert(path, "multiline")
        self.assertIn("Line one\nLine two\nCity",
                      document.tables[0].cell(1, 0).text)

        prose_path = os.path.join(self.root, "prose.pdf")
        prose = canvas.Canvas(prose_path, pagesize=A4)
        for x in (50, 320):
            y = 760
            for i in range(18):
                prose.setFont("Helvetica", 10)
                prose.drawString(
                    x, y,
                    f"paragraph line {i} with ordinary prose text here.")
                y -= 28
        prose.save()
        document = self._convert(prose_path, "prose")
        self.assertEqual(len(document.tables), 0)


if __name__ == "__main__":
    unittest.main()
