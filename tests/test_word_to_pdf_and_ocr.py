import os
import random
import tempfile
import unittest

import pdfplumber
from docx import Document
from docx.enum.text import WD_BREAK

from converter.pdf_to_word import _tables_from_ocr_words
from converter.word_to_pdf import convert_word_to_pdf


def _word(x, y, text, width=50, height=20):
    return (float(x), float(y), float(x + width), float(y + height), text, 95.0)


class OcrTableDetectionTests(unittest.TestCase):
    def test_short_prose_lines_are_not_a_table(self):
        words = []
        for row, y in enumerate((100, 140, 180)):
            words += [_word(60, y, f"satir{row}a"), _word(130, y, "kelime"),
                      _word(230, y, "metin"), _word(320, y, "devam")]
        tables, _ = _tables_from_ocr_words(words)
        self.assertEqual(tables, [])

    def test_full_width_prose_is_not_a_table(self):
        for seed in range(40):
            rnd = random.Random(seed)
            words = []
            for row in range(rnd.randint(4, 10)):
                y = 100 + row * 34
                x = 60.0
                while x < 520:
                    width = rnd.uniform(25, 60)
                    words.append(_word(x, y, "kelime", width=width))
                    x += width + rnd.uniform(6, 12)
            tables, _ = _tables_from_ocr_words(words)
            self.assertEqual(tables, [], f"seed {seed}")

    def test_aligned_grid_is_still_a_table(self):
        words = []
        for row in range(4):
            y = 100 + row * 40
            for col, x in enumerate((60, 180, 300)):
                words.append(_word(x, y, f"r{row}c{col}"))
        tables, used = _tables_from_ocr_words(words)
        self.assertEqual(len(tables), 1)
        self.assertTrue(all(used))


class DocxToPdfPageBreakTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="dc_docx_pdf_tests_")
        self.root = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def test_page_break_moves_following_text_to_next_page(self):
        source = os.path.join(self.root, "break.docx")
        output = os.path.join(self.root, "break.pdf")
        doc = Document()
        doc.add_paragraph("Birinci sayfa metni.")
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("Ikinci sayfa metni.")
        doc.save(source)

        ok, info = convert_word_to_pdf(source, output)
        self.assertTrue(ok, info)
        with pdfplumber.open(output) as pdf:
            self.assertEqual(len(pdf.pages), 2)
            self.assertIn("Birinci", pdf.pages[0].extract_text())
            self.assertIn("Ikinci", pdf.pages[1].extract_text())
            self.assertNotIn("Ikinci", pdf.pages[0].extract_text())


if __name__ == "__main__":
    unittest.main()
