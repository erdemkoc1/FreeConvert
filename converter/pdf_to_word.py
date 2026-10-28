"""PDF -> Word (.docx) dönüştürücü.

pdfminer/pdfplumber ile metin ve tablo, PDFium ile sayfa render, python-docx
ile yerleşik Word çıktısı kullanılır. Word, LibreOffice, OnlyOffice, Java veya
başka bir yan uygulama çağrılmaz. Taranmış sayfalarda gömülü Tesseract
kullanılır; bulut servis kullanılmaz.
"""
from __future__ import annotations

import os
import shutil
import tempfile

_TOL = 2.0
_MIN_GAP = 36.0  # sütun sayılması için en az boşluk (pt)


def _cluster(values: list[float], tol: float = _TOL) -> list[float]:
    vals = sorted(values)
    out: list[float] = []
    for v in vals:
        if out and abs(v - out[-1]) <= tol:
            out[-1] = (out[-1] + v) / 2
        else:
            out.append(v)
    return out


def _pdfplumber_spans(cell_boxes) -> list[tuple[int, int, int, int]]:
    """pdfplumber None hücrelerinden yatay/dikey merge dikdörtgeni üretir."""
    if not cell_boxes:
        return []
    height = len(cell_boxes)
    width = max((len(row) for row in cell_boxes), default=0)
    spans: dict[tuple[int, int], tuple[int, int]] = {}

    for ri, row in enumerate(cell_boxes):
        cj = 0
        while cj < width:
            if cj >= len(row) or row[cj] is not None:
                cj += 1
                continue
            start = cj
            while cj < width and (cj >= len(row) or row[cj] is None):
                cj += 1
            if start > 0 and start - 1 < len(row) and row[start - 1] is not None:
                spans[(ri, start - 1)] = (ri, cj - 1)

    for ri in range(1, height):
        row = cell_boxes[ri]
        for cj in range(width):
            if cj >= len(row) or row[cj] is not None:
                continue
            # Aynı satırdaki yatay None bloğunu tekrar dikey merge sayma.
            if cj > 0 and cj - 1 < len(row) and row[cj - 1] is None:
                continue
            anchor = ri - 1
            while anchor >= 0:
                previous = cell_boxes[anchor]
                if cj < len(previous) and previous[cj] is not None:
                    break
                anchor -= 1
            if anchor >= 0:
                key = (anchor, cj)
                old_end = spans.get(key, (anchor, cj))[1]
                spans[key] = (ri, max(old_end, cj))
    return [(r0, c0, r1, c1) for (r0, c0), (r1, c1) in spans.items()]


def _tables_pdfplumber(plumber_page, words, used):
    """pdfplumber hücre kutularını mevcut Word tablo modeline çevirir.

    Çizgili ve yüksek güvenilirlikli metin hizalı tablo adaylarını dener.
    """
    try:
        from converter.pdfplumber_tables import extract_page_tables
        candidates = extract_page_tables(plumber_page)
    except Exception:
        return []

    tables = []
    for item in candidates:
        grid = item.get("rows") or []
        if len(grid) < 2:
            continue
        C = max(len(row) for row in grid)
        if C < 2:
            continue
        grid = [list(row) + [""] * (C - len(row)) for row in grid]
        R = len(grid)
        gx0, gy0, gx1, gy1 = (float(v) for v in item["bbox"])
        cell_boxes = item.get("cell_boxes") or []
        if len(cell_boxes) != R:
            cell_boxes = []
        cw = (gx1 - gx0) / C
        rh = (gy1 - gy0) / R
        boxes: list[list[tuple[float, float, float, float] | None]] = []
        for ri in range(R):
            row_boxes = []
            source_boxes = cell_boxes[ri] if cell_boxes else []
            for cj in range(C):
                box = source_boxes[cj] if cj < len(source_boxes) else None
                if box is not None and len(box) == 4:
                    row_boxes.append(tuple(float(v) for v in box))
                else:
                    row_boxes.append(None)
            boxes.append(row_boxes)
        spans = _pdfplumber_spans(cell_boxes)

        cells = [[[] for _ in range(C)] for _ in range(R)]
        for wi, word in enumerate(words):
            if used[wi]:
                continue
            cx = (word[0] + word[2]) / 2
            cy = (word[1] + word[3]) / 2
            if not (gx0 - 1 <= cx <= gx1 + 1
                    and gy0 - 1 <= cy <= gy1 + 1):
                continue
            hit = None
            for ri, row_boxes in enumerate(boxes):
                for cj, box in enumerate(row_boxes):
                    if box is None:
                        continue
                    if (box[0] - 1 <= cx <= box[2] + 1
                            and box[1] - 1 <= cy <= box[3] + 1):
                        hit = (ri, cj)
                        break
                if hit:
                    break
            if hit is None:
                # Hizalama toleransı dışında kalan son kelimeleri en yakın
                # mantıklı hücreye koy; tablo metni akışa tekrar sızmasın.
                ri = min(R - 1, max(0, int((cy - gy0) / rh)))
                cj = min(C - 1, max(0, int((cx - gx0) / cw)))
                hit = (ri, cj)
            cells[hit[0]][hit[1]].append(word)
            used[wi] = True

        # pdfplumber hücre metnini daha iyi bilir; kelime eşleşmesi yoksa
        # metni hücre kutusuna çok satırlı sentetik kelimeler olarak koy.
        for ri, row in enumerate(grid):
            for cj, value in enumerate(row):
                value = str(value or "").strip()
                if not value or cells[ri][cj] or boxes[ri][cj] is None:
                    continue
                x0, y0, x1, y1 = boxes[ri][cj]
                line_height = max((y1 - y0) / max(1, value.count("\n") + 1), 8)
                lines = value.splitlines() or [value]
                for li, line in enumerate(lines):
                    cells[ri][cj].append((
                        x0, y0 + li * line_height, x1,
                        y0 + (li + 1) * line_height, line.strip()))

        for row in cells:
            for cell_words in row:
                cell_words.sort(key=lambda word: (word[1], word[0]))
        tables.append((gx0, gy0, gx1, gy1,
                       list(range(R + 1)), list(range(C + 1)), cells,
                       {"spans": spans}))
    return tables


def _detect_tables(plumber_page, words):
    """Yalnızca pdfplumber tablo adaylarını kullanır.

    Dönüş: (tables, rest). Her tablo
    ``(x0, y0, x1, y1, rows, cols, cells, meta)`` biçimindedir;
    ``cells[r][c]`` bir kelime listesi, ``meta`` merge/koordinat bilgisidir.
    """
    used = [False] * len(words)
    tables = _tables_pdfplumber(plumber_page, words, used)
    tables.sort(key=lambda t: t[1])
    rest = [w for wi, w in enumerate(words) if not used[wi]]
    return tables, rest


def _tables_from_ocr_words(words):
    """Taranmış sayfa OCR kelimelerinden güvenilir ızgara benzeri tablo çıkarır."""
    used = [False] * len(words)
    if len(words) < 6:
        return [], used
    normalized = [tuple(word[:5]) for word in words]
    lines = _lines_from_words(normalized)
    # OCR'da bazı hücreler tanınmayabilir; tek kelimelik satırları da
    # aday tabloyu bozmadan ızgara çizgisine dahil et.
    rows = [line for line in lines if line]
    # Üç-dört satırlık düz yazı da sütun gibi görünebilir; gerçek tablo için
    # en az dört satır ızgara kanıtı istenir.
    if len(rows) < 4:
        return [], used

    heights = sorted(max(1.0, w[3] - w[1]) for row in rows for w in row)
    median_height = heights[len(heights) // 2]
    starts = _cluster([w[0] for row in rows for w in row],
                      tol=max(5.0, median_height * 1.4))
    if len(starts) < 2 or len(starts) > 12:
        return [], used

    matrix: list[list[list]] = []
    for line in rows:
        cells: list[list] = [[] for _ in starts]
        for word in line:
            index = min(range(len(starts)),
                        key=lambda i: abs(word[0] - starts[i]))
            cells[index].append(word)
        matrix.append(cells)

    row_count = len(matrix)
    col_count = len(starts)
    nonempty = sum(1 for row in matrix for cell in row if cell)
    if nonempty < max(4, row_count):
        return [], used
    if nonempty / float(row_count * col_count) < 0.25:
        return [], used
    # Gerçek bir ızgarada sütunlar çoğu satırda dolu olur; düz yazı satırlarında
    # kelime başlangıçları rastgele dağılır ve bu koşulu geçemez.
    column_hits = [sum(1 for row in matrix if row[c]) for c in range(col_count)]
    solid_columns = sum(1 for hits in column_hits
                        if hits >= max(2, 0.75 * row_count))
    if solid_columns < 2 or solid_columns < 0.6 * col_count:
        return [], used
    # Hücreler arasında belirgin boşluk olur; kelime arası boşluklu düz yazıda
    # bu boşluk neredeyse hiç görülmez.
    gap_limit = 1.5 * median_height
    gaps = wide_gaps = 0
    for line in rows:
        ordered = sorted(line, key=lambda word: word[0])
        for prev, cur in zip(ordered, ordered[1:]):
            gaps += 1
            if cur[0] - prev[2] >= gap_limit:
                wide_gaps += 1
    if gaps == 0 or wide_gaps < 0.4 * gaps:
        return [], used

    x_min = min(w[0] for row in rows for w in row)
    x_max = max(w[2] for row in rows for w in row)
    x_bounds = [x_min]
    x_bounds.extend((starts[i] + starts[i + 1]) / 2
                    for i in range(len(starts) - 1))
    x_bounds.append(x_max)
    row_bounds: list[tuple[float, float]] = []
    for line in rows:
        y0 = min(w[1] for w in line) - median_height * 0.35
        y1 = max(w[3] for w in line) + median_height * 0.35
        row_bounds.append((y0, y1))
    cell_boxes = []
    for ri, row in enumerate(matrix):
        y0, y1 = row_bounds[ri]
        cell_boxes.append([
            (x_bounds[cj], y0, x_bounds[cj + 1], y1)
            for cj in range(col_count)
        ])

    word_ids = {id(word): index for index, word in enumerate(normalized)}
    for row in matrix:
        for cell in row:
            for word in cell:
                used[word_ids[id(word)]] = True
            cell.sort(key=lambda word: (word[1], word[0]))

    gx0, gy0 = x_bounds[0], row_bounds[0][0]
    gx1, gy1 = x_bounds[-1], row_bounds[-1][1]
    table = (gx0, gy0, gx1, gy1,
             [row_bounds[0][0]] + [b[1] for b in row_bounds],
             x_bounds, matrix, {"cell_boxes": cell_boxes, "ocr": True})
    return [table], used


def _split_columns(words, page_width: float):
    if not words:
        return []
    BIN = 5.0
    nbins = max(int(page_width / BIN) + 1, 1)
    covered = [False] * nbins
    for w in words:
        a = max(int(w[0] / BIN), 0)
        b = min(int(w[2] / BIN), nbins - 1)
        for i in range(a, b + 1):
            covered[i] = True
    bounds = [0.0]
    start = None
    for i, c in enumerate(covered):
        if not c and start is None:
            start = i
        elif c and start is not None:
            g0, g1 = start * BIN, i * BIN
            if g1 - g0 >= _MIN_GAP and g0 > 20 and g1 < page_width - 20:
                bounds.append((g0 + g1) / 2)
            start = None
    if start is not None:
        g0, g1 = start * BIN, page_width
        if g1 - g0 >= _MIN_GAP and g0 > 20 and g1 < page_width - 20:
            bounds.append((g0 + g1) / 2)
    bounds.append(page_width)
    cols: list[list] = [[] for _ in range(len(bounds) - 1)]
    for w in words:
        cx = (w[0] + w[2]) / 2
        idx = len(cols) - 1
        for k in range(len(bounds) - 1):
            if bounds[k] <= cx <= bounds[k + 1]:
                idx = k
                break
        cols[idx].append(w)
    return [c for c in cols if c]


def _lines_from_words(words):
    """Kelimeleri satırlara grupla (merkez-y yakınlığı).

    Kelime kutuları satırlar arası boşluğu da kapsayabildiği için
    kutu-örtüşmesi yerine satır merkezleri kullanılır.
    """
    if not words:
        return []
    hs = sorted(w[3] - w[1] for w in words)
    thr = max(hs[len(hs) // 2] * 0.35, 2.0)
    seq = sorted(words, key=lambda v: ((v[1] + v[3]) / 2, v[0]))
    lines: list[list] = []
    avgs: list[float] = []
    for w in seq:
        wc = (w[1] + w[3]) / 2
        if lines and abs(wc - avgs[-1]) <= thr:
            lines[-1].append(w)
            avgs[-1] = (avgs[-1] * (len(lines[-1]) - 1) + wc) / len(lines[-1])
        else:
            lines.append([w])
            avgs.append(wc)
    for ln in lines:
        ln.sort(key=lambda v: v[0])
    return lines


def _line_style(words, spans):
    """Satırdaki kelimelerin ortak stili: (all_bold, all_italic)."""
    bolds, italics = [], []
    for w in words:
        cx, cy = (w[0] + w[2]) / 2, (w[1] + w[3]) / 2
        hit = [s for s in spans
               if s[0] - 1 <= cx <= s[2] + 1 and s[1] - 1 <= cy <= s[3] + 1]
        if not hit:
            return False, False
        bolds.append(any(s[4] for s in hit))
        italics.append(any(s[5] for s in hit))
    if not bolds:
        return False, False
    return all(bolds), all(italics)


def _set_page_geometry(section, rect) -> None:
    """PDF sayfa boyutunu Word bölümüne taşır."""
    try:
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt
        if hasattr(rect, "width"):
            width, height = float(rect.width), float(rect.height)
        else:
            width, height = float(rect[0]), float(rect[1])
        section.orientation = (WD_ORIENT.LANDSCAPE if width > height
                              else WD_ORIENT.PORTRAIT)
        section.page_width = Pt(width)
        section.page_height = Pt(height)
        section.left_margin = Pt(36)
        section.right_margin = Pt(36)
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
    except Exception:
        pass


def _emit_table(doc, cells, spans, meta=None) -> None:
    """Hücre kelime listelerinden çok satırlı Word tablosu kurar."""
    if not cells or not cells[0]:
        return
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = doc.add_table(rows=len(cells), cols=len(cells[0]))
    try:
        tbl.style = "Table Grid"
        tbl.autofit = True
    except Exception:
        pass

    merge_spans = list((meta or {}).get("spans", []))
    covered: set[tuple[int, int]] = set()
    for span in merge_spans:
        try:
            r0, c0, r1, c1 = (int(v) for v in span)
            if r0 >= r1 and c0 >= c1:
                continue
            tbl.cell(r0, c0).merge(tbl.cell(r1, c1))
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    if (rr, cc) != (r0, c0):
                        covered.add((rr, cc))
        except Exception:
            continue

    # PDF hücre genişliklerini Word'e yakın ölçekte taşı.
    cell_boxes = list((meta or {}).get("cell_boxes", []))
    if cell_boxes:
        try:
            widths = []
            for cj in range(len(cells[0])):
                widths.append(max((box[cj][2] - box[cj][0]
                                   for box in cell_boxes
                                   if cj < len(box) and box[cj] is not None),
                                  default=0))
            total = sum(widths)
            if total > 0:
                scale = min(1.0, (6.5 * 72.0) / total)
                tbl.autofit = False
                for cj, width in enumerate(widths):
                    if width <= 0:
                        continue
                    tbl.columns[cj].width = Pt(width * scale)
                    for row in tbl.rows:
                        row.cells[cj].width = Pt(width * scale)
        except Exception:
            pass

    # Uzun tablolarda başlık satırı sonraki sayfalarda tekrarlansın.
    try:
        tr_pr = tbl.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)
    except Exception:
        pass

    for ri, row in enumerate(cells):
        try:
            tr_pr = tbl.rows[ri]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
        except Exception:
            pass
        for cj, cw in enumerate(row):
            if cj >= len(cells[0]) or (ri, cj) in covered:
                continue
            cell = tbl.cell(ri, cj)
            cell.text = ""
            visual_lines = _lines_from_words(cw) if cw else []
            if not visual_lines and cw:
                visual_lines = [cw]
            first_paragraph = True
            for line_words in visual_lines:
                raw = " ".join(str(v[4]) for v in line_words).strip()
                for part in (raw.splitlines() or [raw]):
                    part = part.strip()
                    if not part:
                        continue
                    paragraph = (cell.paragraphs[0] if first_paragraph
                                 else cell.add_paragraph())
                    first_paragraph = False
                    try:
                        paragraph.paragraph_format.space_after = Pt(0)
                    except Exception:
                        pass
                    bold, italic = _line_style(line_words, spans)
                    run = paragraph.add_run(part)
                    run.bold = bold
                    run.italic = italic


def _words_from_plumber_page(page) -> list[tuple]:
    try:
        extracted = page.extract_words(
            x_tolerance=2, y_tolerance=3, keep_blank_chars=False,
            use_text_flow=False)
    except Exception:
        return []
    words = []
    for item in extracted:
        text = str(item.get("text", "")).strip()
        if text:
            words.append((float(item["x0"]), float(item["top"]),
                          float(item["x1"]), float(item["bottom"]), text))
    return words


def convert_pdf_to_word(input_path: str, output_path: str,
                        mode: str = "adaptif") -> tuple[bool, str]:
    """PDF'i Word'e çevirir; ``mode`` geriye uyum için yok sayılır.

    pdfplumber/pdfminer metin ve tablo çıkarır, PDFium sayfa render ve
    görsel çıkarımı yapar, python-docx yerleşik çıktı üretir. Hiçbir yan
    uygulama çağrılmaz.
    """
    ext = os.path.splitext(input_path)[1].lower()
    if ext != ".pdf":
        return False, "Sadece .pdf destekleniyor."
    os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)
    try:
        if os.path.exists(output_path):
            os.remove(output_path)
    except Exception:
        pass

    try:
        import pdfplumber
        import pypdfium2 as pdfium
        from pypdfium2 import raw as pdfium_raw
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Pt, Inches
    except Exception as exc:
        return False, f"Gerekli açık kaynak kütüphaneler yüklenemedi: {exc}"

    plumber_doc = None
    pdfium_doc = None
    tmpdir = tempfile.mkdtemp(prefix="pdf2word_")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    n_tables = 0
    n_ocr = 0
    try:
        plumber_doc = pdfplumber.open(input_path)
        pdfium_doc = pdfium.PdfDocument(input_path)
        if len(plumber_doc.pages) == 0:
            return False, "PDF içinde sayfa bulunamadı."
        _set_page_geometry(
            doc.sections[0],
            (float(plumber_doc.pages[0].width),
             float(plumber_doc.pages[0].height)))

        for page_no, plumber_page in enumerate(plumber_doc.pages):
            pdfium_page = pdfium_doc[page_no]
            words = _words_from_plumber_page(plumber_page)
            spans = []

            ocr_text = ""
            ocr_used = False
            if not words:
                try:
                    from converter.ocr import ocr_pdf_page_words
                    ocr_words = ocr_pdf_page_words(pdfium_page)
                    if len(ocr_words) >= 6:
                        words = [tuple(word[:5]) for word in ocr_words]
                        ocr_used = True
                except Exception:
                    ocr_words = []
                if not ocr_used:
                    try:
                        from converter.ocr import ocr_pdf_page_text
                        ocr_text = ocr_pdf_page_text(pdfium_page)
                    except Exception:
                        ocr_text = ""

            tables, rest = _detect_tables(plumber_page, words)
            if ocr_used and not tables:
                tables, ocr_used_flags = _tables_from_ocr_words(words)
                if tables:
                    rest = [word for index, word in enumerate(words)
                            if not ocr_used_flags[index]]
            if ocr_used:
                n_ocr += 1
            page_width = float(plumber_page.width)
            columns = (_split_columns(rest, page_width)
                       or ([rest] if rest else []))

            col_tables: list[list] = [[] for _ in columns]
            orphans: list = []
            for table in tables:
                if not columns:
                    orphans.append(table)
                    continue
                center_x = (table[0] + table[2]) / 2
                index = 0
                low = 0.0
                for column_index, column in enumerate(columns):
                    high = max((word[2] for word in column),
                               default=page_width)
                    if low <= center_x <= high:
                        index = column_index
                        break
                    low = high
                col_tables[index].append(table)

            for column_index, column in enumerate(columns):
                events = []
                for line in _lines_from_words(column):
                    events.append((min(word[1] for word in line),
                                   "line", line))
                for table in col_tables[column_index]:
                    events.append((table[1], "table", table))
                events.sort(key=lambda event: event[0])
                for _, kind, payload in events:
                    if kind == "line":
                        text = " ".join(word[4] for word in payload).strip()
                        if text:
                            paragraph = doc.add_paragraph()
                            paragraph.add_run(text)
                    else:
                        _, _, _, _, _, _, cells, meta = payload
                        _emit_table(doc, cells, spans, meta)
                        n_tables += 1
            if not columns:
                for table in orphans:
                    _, _, _, _, _, _, cells, meta = table
                    _emit_table(doc, cells, spans, meta)
                    n_tables += 1

            if ocr_text.strip():
                for paragraph_text in ocr_text.splitlines():
                    paragraph_text = paragraph_text.strip()
                    if paragraph_text:
                        doc.add_paragraph(paragraph_text)
                if not ocr_used:
                    n_ocr += 1

            # OCR tabloyu oluşturduysa tam sayfa rasterını tekrar ekleme.
            if not (ocr_used and tables):
                try:
                    image_type = pdfium_raw.FPDF_PAGEOBJ_IMAGE
                    for image_number, obj in enumerate(
                            pdfium_page.get_objects()):
                        if getattr(obj, "type", None) != image_type:
                            continue
                        image_path = os.path.join(
                            tmpdir, f"p{page_no}_{image_number}.bin")
                        try:
                            obj.extract(image_path)
                            doc.add_picture(image_path, width=Inches(5.5))
                        except Exception:
                            continue
                        finally:
                            try:
                                obj.close()
                            except Exception:
                                pass
                except Exception:
                    pass

            if page_no < len(plumber_doc.pages) - 1:
                next_section = doc.add_section(WD_SECTION.NEW_PAGE)
                _set_page_geometry(
                    next_section,
                    (float(plumber_doc.pages[page_no + 1].width),
                     float(plumber_doc.pages[page_no + 1].height)))

        doc.save(output_path)
        info = f"{n_tables} tablo, {len(doc.paragraphs)} paragraf aktarıldı."
        if n_ocr:
            info += f" ({n_ocr} sayfa OCR)"
        return True, info + " Dahili offline motor."
    except Exception as exc:
        return False, f"Dönüştürme hatası: {exc}"
    finally:
        for resource in (plumber_doc, pdfium_doc):
            try:
                if resource is not None:
                    resource.close()
            except Exception:
                pass
        shutil.rmtree(tmpdir, ignore_errors=True)
