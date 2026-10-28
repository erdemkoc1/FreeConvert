"""Dönüştürücülerin ortak, tamamen yerel yardımcıları.

Uygulama dışında Word/LibreOffice/OnlyOffice aranmaz. PDF metin ve tablo
işlemleri pdfminer/pdfplumber, sayfa render işlemleri PDFium tarafından
yapılır.
"""
from __future__ import annotations

import os
import shutil
import sys


def bundled_tesseract_dir() -> str | None:
    """Exe içine gömülü Tesseract runtime klasörü (tesseract_runtime/).

    Paketlenmiş (frozen) derlemede sys._MEIPASS altında, geliştirme
    modunda proje kökü altında arar. Bulamazsa None.
    """
    exe = "tesseract.exe" if os.name == "nt" else "tesseract"
    bases = []
    if getattr(sys, "frozen", False):
        bases.append(getattr(sys, "_MEIPASS", "") or "")
    bases.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    for b in bases:
        if not b:
            continue
        d = os.path.join(b, "tesseract_runtime")
        if os.path.exists(os.path.join(d, exe)):
            return d
    return None


def tessdata_dir() -> str | None:
    """Gömülü dil verisi klasörü (tur/eng). Bulamazsa None."""
    d = bundled_tesseract_dir()
    if d and os.path.exists(os.path.join(d, "tessdata")):
        return os.path.join(d, "tessdata")
    return None


def docx_has_text(path: str) -> bool:
    """Docx açılıp içinde gerçek metin var mı? Boş çıktıları ele."""
    try:
        from docx import Document
        doc = Document(path)
    except Exception:
        return False
    try:
        for p in doc.paragraphs:
            if p.text.strip():
                return True
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        return True
    except Exception:
        return False
    return False


def pdf_has_content(path: str) -> bool:
    """PDF açılıp sayfa + (metin, görsel veya çizim) var mı?"""
    try:
        import pdfplumber
        with pdfplumber.open(path) as doc:
            if not doc.pages:
                return False
            for page in doc.pages:
                try:
                    if page.extract_text(x_tolerance=2, y_tolerance=3).strip():
                        return True
                    if page.images or page.lines or page.rects or page.curves:
                        return True
                except Exception:
                    continue
        return False
    except Exception:
        pass

    # pdfplumber yoksa permissive PDFium metin katmanına bak.
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(path)
        try:
            if len(doc) == 0:
                return False
            for index in range(len(doc)):
                page = doc[index]
                text_page = page.get_textpage()
                try:
                    if text_page.get_text_range().strip():
                        return True
                finally:
                    text_page.close()
                    page.close()
        finally:
            doc.close()
    except Exception:
        return False
    return False


def find_tesseract() -> str | None:
    """Yerel OCR motorunu bul.

    Paketlenmiş sürümde yalnızca exe içine gömülü motor kabul edilir; başka
    bilgisayarda kurulu sistem Tesseract'ına sessizce bağımlı olunmaz.
    Geliştirme çalıştırırken sistem kurulumu kullanılabilir.
    """
    d = bundled_tesseract_dir()
    if d:
        p = os.path.join(d, "tesseract.exe" if os.name == "nt" else "tesseract")
        if os.path.exists(p):
            return p
    if getattr(sys, "frozen", False):
        return None
    found = shutil.which("tesseract")
    if found:
        return found
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.join(os.environ.get("LOCALAPPDATA", ""), r"Tesseract-OCR\tesseract.exe"),
        "/usr/bin/tesseract",
        "/usr/local/bin/tesseract",
    ]
    for c in candidates:
        if c and os.path.exists(c):
            return c
    return None
